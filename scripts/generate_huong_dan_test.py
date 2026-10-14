"""Script to generate huong_dan_test.docx with standardized Times New Roman font and readable layout."""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_font_times(run, font_name="Times New Roman", size_pt=11, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" '
        f'w:cs="{font_name}" w:eastAsia="{font_name}"/>'
    )
    rPr.append(rFonts)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="6"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout(doc, text, title="LƯU Ý QUAN TRỌNG", border_color="005A9C", bg_color="F4F8FA"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=220, right=220)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="28" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r_title = p.add_run(f"📌 {title}: ")
    set_font_times(r_title, size_pt=11, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    r_txt = p.add_run(text)
    set_font_times(r_txt, size_pt=11, color_rgb=RGBColor(0x33, 0x33, 0x33))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def build_test_guide(output_path):
    doc = docx.Document()
    
    # Set default style font
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.space_after = Pt(6)
    
    # Page setup: Margins 1 inch (2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    # ─── HEADER TITLE BLOCK ───
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_inst.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO — TRƯỜNG ĐẠI HỌC CNTT & TT VIỆT - HÀN (VKU)\n")
    set_font_times(r1, size_pt=11.5, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    r2 = p_inst.add_run("KHOA KỸ THUẬT MÁY TÍNH VÀ ĐIỆN TỬ\n")
    set_font_times(r2, size_pt=11.5, bold=True, color_rgb=RGBColor(0x55, 0x55, 0x55))
    
    r_div = p_inst.add_run("━━━━━━━━━━━━━━━━━━━ ◆ ━━━━━━━━━━━━━━━━━━━\n\n")
    set_font_times(r_div, size_pt=10, color_rgb=RGBColor(0x88, 0x88, 0x88))
    
    p_main = doc.add_paragraph()
    p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_main.paragraph_format.line_spacing = 1.2
    
    r_title = p_main.add_run("TÀI LIỆU HƯỚNG DẪN KIỂM THỬ TOÀN DIỆN\n")
    set_font_times(r_title, size_pt=17, bold=True, color_rgb=RGBColor(0xB2, 0x22, 0x22))
    
    r_sub = p_main.add_run("HỆ THỐNG GIÁM SÁT VÀ CẢNH BÁO SỚM CHÁY & KHÓI BẰNG TRÍ TUỆ NHÂN TẠO\n(FIREEYES AI — SYSTEM TEST GUIDE)\n\n")
    set_font_times(r_sub, size_pt=13, bold=True, color_rgb=RGBColor(0x1B, 0x36, 0x5D))
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.line_spacing = 1.2
    r_meta = p_meta.add_run(
        "Nhóm thực hiện (Nhóm 4 thành viên):\n"
        "• Nguyễn Đình Anh — Trưởng nhóm (Phụ trách Backend API, MVC Controller & Quản trị CLI)\n"
        "• Phú Anh — Thành viên chính (Phụ trách Kiến trúc MVC, Frontend SPA & Real-time Vision)\n"
        "• Huy Hoàng — Thành viên (Phụ trách Thiết kế CSDL SQLite v2 & Toàn vẹn dữ liệu)\n"
        "• Thùy Ngân — Thành viên (Phụ trách Mô hình AI YOLO11n & Đánh giá thị giác máy tính)\n\n"
        "Môi trường kiểm thử: Windows 10/11 (64-bit) | Python 3.10+ | FastAPI | SQLite v2\n"
        "Phiên bản tài liệu: v2.5 (Cập nhật tháng 10/2026)\n"
    )
    set_font_times(r_meta, size_pt=10.5, italic=True, color_rgb=RGBColor(0x44, 0x44, 0x44))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # ─── SECTION 1: MỤC TIÊU VÀ PHẠM VI ───
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True
    r_h1 = h1.add_run("1. MỤC TIÊU VÀ PHẠM VI TÀI LIỆU KIỂM THỬ")
    set_font_times(r_h1, size_pt=13.5, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.line_spacing = 1.25
    p_desc.paragraph_format.space_after = Pt(6)
    p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p_desc.add_run(
        "Tài liệu này được biên soạn nhằm chuẩn hóa toàn bộ quy trình kiểm thử hệ thống FireEyes AI từ khâu chuẩn bị môi trường, "
        "khởi động máy chủ, kiểm thử chức năng giao diện người dùng (Functional Tests) đến kiểm thử hồi quy tự động ở tầng CSDL và API (Automated Unit Tests). "
        "Mục tiêu là hỗ trợ giảng viên hướng dẫn, hội đồng đánh giá và các thành viên độc lập tái lập chính xác mọi ca kiểm thử: "
        "nhận diện ngọn lửa trực tiếp bằng bật lửa, thuật toán lọc da người chống báo cháy giả, tải ảnh hiện trường, bật/tắt camera phần cứng, "
        "tạo chuỗi dữ liệu giả lập và quy trình xử lý đóng sự cố cảnh báo."
    )
    set_font_times(r, size_pt=11)
    
    add_callout(
        doc,
        "Trước khi tiến hành kiểm thử, hãy đảm bảo webcam của bạn đang hoạt động bình thường, đã cấp quyền Camera trên trình duyệt và loa máy tính bật mức âm lượng vừa phải để nghe còi hú báo động.",
        title="YÊU CẦU TIÊN QUYẾT"
    )
    
    # ─── SECTION 2: MÔI TRƯỜNG & KHỞI ĐỘNG ───
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    r_h2 = h2.add_run("2. CHUẨN BỊ MÔI TRƯỜNG VÀ QUY TRÌNH KHỞI ĐỘNG")
    set_font_times(r_h2, size_pt=13.5, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    tbl_env = doc.add_table(rows=6, cols=3)
    tbl_env.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_env)
    
    headers = ["Thành phần", "Cấu hình khuyến nghị", "Ghi chú & Trạng thái"]
    for i, title in enumerate(headers):
        cell = tbl_env.cell(0, i)
        cell.text = title
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        set_font_times(p.runs[0], size_pt=10.5, bold=True, color_rgb=RGBColor(0xFF, 0xFF, 0xFF))
        
    env_rows = [
        ("Hệ điều hành", "Windows 10 / 11 (64-bit)", "Đã kiểm nghiệm tương thích 100% trên Windows"),
        ("Môi trường Python", "Python 3.10+ (Đã cấu hình tại .venv)", "Sử dụng đầy đủ thư viện trong thư mục .venv"),
        ("Máy chủ Web / API", "FastAPI + Uvicorn (Port 8000)", "Phục vụ đồng thời REST API và Web tĩnh SPA"),
        ("Cơ sở dữ liệu", "SQLite 3 (File: backend/data/fireeyes.db)", "Lược đồ v2 gồm đầy đủ 11 bảng nghiệp vụ chuẩn"),
        ("Thiết bị ngoại vi", "Webcam tích hợp / USB, Bật lửa ga mini, Loa", "Sử dụng trực tiếp cho các ca kiểm nghiệm thực tế"),
    ]
    
    for row_idx, data in enumerate(env_rows, 1):
        bg = "F7F9FC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = tbl_env.cell(row_idx, col_idx)
            cell.text = text
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
            p = cell.paragraphs[0]
            set_font_times(p.runs[0], size_pt=10.5, color_rgb=RGBColor(0x33, 0x33, 0x33))
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    
    # Quy trình khởi động
    p_step_title = doc.add_paragraph()
    p_step_title.paragraph_format.space_before = Pt(6)
    p_step_title.paragraph_format.space_after = Pt(4)
    r = p_step_title.add_run("Quy trình 4 bước khởi động hệ thống:")
    set_font_times(r, size_pt=11.5, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    steps = [
        ("Bước 1: Mở Terminal PowerShell", "Mở Windows PowerShell tại thư mục dự án: C:\\Users\\ADMIN\\Desktop\\fire_detection_ai-main."),
        ("Bước 2: Kích hoạt máy chủ FastAPI", "Thực thi lệnh khởi chạy backend và giao diện web:\n.venv\\Scripts\\python.exe -m uvicorn backend.api:app --host 127.0.0.1 --port 8000 --reload"),
        ("Bước 3: Kiểm tra trạng thái máy chủ (Health Check)", "Mở trình duyệt truy cập: http://127.0.0.1:8000/health — Kết quả trả về JSON: {\"status\": \"ok\"}."),
        ("Bước 4: Truy cập Giao diện Web SPA", "Mở trình duyệt truy cập: http://127.0.0.1:8000/ để bắt đầu sử dụng đầy đủ các tính năng."),
    ]
    
    for st_title, st_desc in steps:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(3)
        r_num = p.add_run(f"• {st_title}: ")
        set_font_times(r_num, size_pt=11, bold=True, color_rgb=RGBColor(0x00, 0x5A, 0x9C))
        
        r_txt = p.add_run(st_desc)
        set_font_times(r_txt, size_pt=10.5)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # ─── SECTION 3: DANH MỤC CÁC CA KIỂM THỬ (TEST CASES) ───
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    r_h3 = h3.add_run("3. MA TRẬN VÀ KỊCH BẢN KIỂM THỬ CHI TIẾT (TEST CASES)")
    set_font_times(r_h3, size_pt=13.5, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    test_cases = [
        {
            "id": "TC-01",
            "name": "Kiểm tra Bật / Tắt Webcam và Giải phóng Phần cứng Camera",
            "objective": "Đảm bảo rằng khi người dùng bấm nút 'Tắt camera', webcam phần cứng ngắt kết nối hoàn toàn (tắt đèn LED camera trên laptop), dừng luồng video, và canvas hiển thị giao diện offline chuẩn.",
            "steps": [
                "1. Truy cập giao diện web tại mục 'Giám sát camera AI'.",
                "2. Kiểm tra webcam đang hoạt động (đèn webcam sáng, video hiển thị).",
                "3. Nhấp vào nút 'Tắt camera' (màu đỏ).",
                "4. Quan sát đèn phần cứng webcam và khung hình canvas.",
                "5. Nhấp nút 'Bật camera' để kích hoạt trở lại."
            ],
            "expected": "Khi tắt: Đèn LED webcam tắt hẳn, luồng video MediaStream ngắt hoàn toàn (tất cả tracks dừng), màn hình canvas hiển thị biển báo màu tối 'CAMERA ĐÃ TẮT (STOPPED)', xóa bỏ các bounding box cũ. Khi bật lại: Webcam khởi động lại trơn tru trong 1-2 giây.",
            "status": "PASS (Đạt 100%)"
        },
        {
            "id": "TC-02",
            "name": "Kiểm tra Nhận Diện Ngọn Lửa Thực Tế Bằng Bật Lửa (Lighter Test)",
            "objective": "Kiểm tra khả năng nhận diện ngọn lửa trực tiếp từ bật lửa ga/que diêm theo thời gian thực và kích hoạt báo động khẩn cấp.",
            "steps": [
                "1. Đảm bảo camera AI đang ở trạng thái 'ĐANG GIÁM SÁT'.",
                "2. Sử dụng bật lửa ga cầm tay, bật lửa ở khoảng cách 30cm - 60cm phía trước camera.",
                "3. Giữ ngọn lửa cháy ổn định trong khoảng 1.5 đến 2 giây (để thuật toán cửa sổ trượt 3/5 khung hình xác nhận).",
                "4. Quan sát phản ứng trên màn hình canvas, âm thanh và bảng sự cố."
            ],
            "expected": "1. Khung chữ nhật Bounding Box màu đỏ xuất hiện bao quanh ngọn lửa với nhãn 'LỬA (FIRE) 90+%'.\n2. Cửa sổ trượt xác nhận đạt điều kiện 3/5 frames liên tiếp.\n3. Còi hú báo động âm thanh Web Audio kêu to dồn dập.\n4. Banner cảnh báo khẩn cấp màu đỏ nhấp nháy 'PHÁT HIỆN CHÁY TỨC THÌ!'.\n5. Sự kiện mới được tự động ghi nhận vào CSDL, bảng 'Nhật ký sự kiện' cập nhật dòng mới kèm ảnh chụp.",
            "status": "PASS (Đạt 100%)"
        },
        {
            "id": "TC-03",
            "name": "Kiểm tra Thuật Toán Lọc Chống Báo Động Giả (Anti-False Alarm / Skin Filter)",
            "objective": "Đảm bảo hệ thống KHÔNG báo cháy giả khi đưa bàn tay, khuôn mặt người hoặc ánh đèn phòng/tường xám vào trước camera.",
            "steps": [
                "1. Đang mở camera giám sát phòng bình thường (chưa có lửa).",
                "2. Đưa khuôn mặt lại gần camera, vẫy bàn tay qua lại trước ống kính.",
                "3. Chiếu ánh đèn pin điện thoại hoặc ánh đèn vàng vào góc camera.",
                "4. Quan sát trong 30 giây."
            ],
            "expected": "Hệ thống áp dụng bộ lọc vật lý màu ngọn lửa (độ bão hòa S >= 0.58, độ chói >= 630). Vùng da người (S thuộc khoảng [0.25, 0.45]) và tường tĩnh bị loại trừ hoàn toàn. KHÔNG xuất hiện bounding box lửa trên khuôn mặt/bàn tay, KHÔNG phát còi báo động, trạng thái camera duy trì 'AN TOÀN - KHÔNG PHÁT HIỆN CHÁY KHÓI'.",
            "status": "PASS (Đạt 100%)"
        },
        {
            "id": "TC-04",
            "name": "Kiểm tra Chức Năng Phân Tích Ảnh Tải Lên (Static Image Inference)",
            "objective": "Kiểm tra tính năng tải lên ảnh tĩnh chụp hiện trường cháy/khói để phân tích và đánh dấu bounding box.",
            "steps": [
                "1. Tại tab 'Tải ảnh phân tích', kéo thả hoặc nhấn chọn một file ảnh JPG/PNG có đám cháy hoặc khói.",
                "2. Quan sát khung hiển thị kết quả phân tích.",
                "3. Kiểm tra danh sách phát hiện và bảng sự cố."
            ],
            "expected": "Ảnh được hiển thị rõ nét, thuật toán quét điểm ảnh và vẽ bounding box chính xác bao quanh đám cháy/khói, hiển thị điểm tin cậy Confidence, tự động lưu kết quả vào CSDL và cập nhật thẻ số liệu.",
            "status": "PASS (Đạt 100%)"
        },
        {
            "id": "TC-05",
            "name": "Kiểm tra Chức Năng Tạo Dữ Liệu Giả Lập Demo (MVC Seed Demo Test)",
            "objective": "Kiểm tra nút 'Tạo demo' trên thanh công cụ giao diện và lệnh CLI tạo chuỗi dữ liệu giả lập chuẩn theo mô hình MVC.",
            "steps": [
                "1. Nhấp nút 'Tạo demo' (nút màu xanh lá hoặc tím trên thanh công cụ).",
                "2. Quan sát thông báo popup Toast và bảng điều khiển thống kê.",
                "3. Kiểm tra dòng sự cố demo vừa được tạo."
            ],
            "expected": "Tạo thành công bản ghi khép kín: Frame -> Detection -> Incident -> Simulated Notification. Hệ thống tự động re-enable camera nếu đang bị tắt, không xảy ra lỗi CameraUnavailableError. Các thẻ thống kê tăng đúng 1 đơn vị.",
            "status": "PASS (Đạt 100%)"
        },
        {
            "id": "TC-06",
            "name": "Kiểm tra Quy Trình Tiếp Nhận & Xử Lý Sự Cố (Acknowledge & Resolve)",
            "objective": "Kiểm tra tính năng cán bộ trực ban xác nhận đã xem cảnh báo và đóng sự kiện (đã dập lửa hoặc báo động giả).",
            "steps": [
                "1. Tại một sự kiện đang mở (Status: Open), nhấn nút 'Xác nhận'.",
                "2. Nhập tên cán bộ trực ban (ví dụ: 'Trực ban PCCC VKU').",
                "3. Nhấn nút 'Xử lý' -> Chọn trạng thái 'Đã giải quyết' (Resolved) hoặc 'Báo động giả' (False alarm).",
                "4. Nhập ghi chú giải quyết: 'Đã dùng bình chữa cháy CO2 dập tắt lúc 06:40' và bấm Lưu."
            ],
            "expected": "Trạng thái sự cố chuyển sang badge xanh 'Đã xử lý' hoặc vàng 'Báo động giả'. Còi hú ngắt hoàn toàn. CSDL cập nhật chính xác thời điểm closed_at và trường resolution_note.",
            "status": "PASS (Đạt 100%)"
        },
        {
            "id": "TC-07",
            "name": "Kiểm tra Modal Xem Ảnh Bằng Chứng (Snapshot Preview & Cache)",
            "objective": "Kiểm tra khả năng mở modal phóng to ảnh chụp bằng chứng thời điểm phát hiện cháy.",
            "steps": [
                "1. Tìm một sự cố có ảnh bằng chứng trong bảng nhật ký.",
                "2. Nhấp vào nút/icon xem ảnh snapshot.",
                "3. Quan sát modal hiển thị."
            ],
            "expected": "Modal phóng to mở ra mượt mà, hiển thị ảnh bằng chứng với độ phân giải cao kèm bounding box ngọn lửa, thông tin camera, thời gian chụp UTC và nút tải ảnh về máy.",
            "status": "PASS (Đạt 100%)"
        },
        {
            "id": "TC-08",
            "name": "Kiểm tra Các Lệnh Dòng Lệnh Quản Trị Backend (Backend CLI Tools)",
            "objective": "Kiểm tra 4 lệnh dòng lệnh CLI độc lập: init, stats, incidents, demo.",
            "steps": [
                "Chạy lần lượt trong terminal:\n1. .venv\\Scripts\\python.exe -m backend.main init\n2. .venv\\Scripts\\python.exe -m backend.main stats\n3. .venv\\Scripts\\python.exe -m backend.main incidents\n4. .venv\\Scripts\\python.exe -m backend.main demo"
            ],
            "expected": "Cả 4 lệnh thực thi trơn tru, mã thoát returncode = 0. Đầu ra JSON tiếng Việt có dấu in đúng định dạng UTF-8, không bị lỗi UnicodeEncodeError charmap trên console Windows.",
            "status": "PASS (Đạt 100%)"
        },
        {
            "id": "TC-09",
            "name": "Kiểm tra Toàn Diện Bộ Unit Tests Tự Động (Automated Regression Tests)",
            "objective": "Kiểm tra 100% các ràng buộc toàn vẹn, giao dịch ACID, trigger, API và Controller qua bộ test tự động.",
            "steps": [
                "Chạy lệnh kiểm thử tự động:\n.venv\\Scripts\\python.exe -m unittest discover -s backend/tests -v"
            ],
            "expected": "Toàn bộ 26/26 bài test trong test_api.py và test_database.py đều đạt (Ran 26 tests - OK), thời gian thực thi < 3 giây.",
            "status": "PASS (26/26 OK)"
        },
        {
            "id": "TC-10",
            "name": "Kiểm tra Tài Liệu Hóa và Thử Nghiệm API Swagger UI (OpenAPI)",
            "objective": "Kiểm tra tính sẵn sàng của hệ thống API RESTful phục vụ kết nối bên ngoài.",
            "steps": [
                "1. Truy cập trình duyệt tại: http://127.0.0.1:8000/docs",
                "2. Kiểm tra danh mục các endpoints: Locations, Cameras, Datasets, Models, Frames, Alerts, Contacts, Notifications.",
                "3. Thử nghiệm gọi GET /api/stats trực tiếp trên Swagger UI."
            ],
            "expected": "Giao diện Swagger tương tác đầy đủ, phản hồi mã 200 OK kèm dữ liệu JSON chính xác.",
            "status": "PASS (Đạt 100%)"
        },
    ]
    
    for tc in test_cases:
        p_tc = doc.add_paragraph()
        p_tc.paragraph_format.space_before = Pt(10)
        p_tc.paragraph_format.space_after = Pt(2)
        p_tc.paragraph_format.keep_with_next = True
        
        r_id = p_tc.add_run(f"[{tc['id']}] {tc['name']}")
        set_font_times(r_id, size_pt=12, bold=True, color_rgb=RGBColor(0xB2, 0x22, 0x22))
        
        tbl_tc = doc.add_table(rows=4, cols=2)
        tbl_tc.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl_tc)
        
        col_widths = [Inches(1.8), Inches(4.7)]
        
        tc_data = [
            ("Mục tiêu kiểm thử", tc["objective"]),
            ("Các bước thực hiện", "\n".join(tc["steps"]) if isinstance(tc["steps"], list) else tc["steps"]),
            ("Kết quả mong đợi", tc["expected"]),
            ("Đánh giá kết quả", tc["status"]),
        ]
        
        for r_idx, (lbl, val) in enumerate(tc_data):
            cell_lbl = tbl_tc.cell(r_idx, 0)
            cell_val = tbl_tc.cell(r_idx, 1)
            
            cell_lbl.width = col_widths[0]
            cell_val.width = col_widths[1]
            
            cell_lbl.text = lbl
            cell_val.text = val
            
            set_cell_background(cell_lbl, "F4F8FA")
            set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
            set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)
            
            p_l = cell_lbl.paragraphs[0]
            set_font_times(p_l.runs[0], size_pt=10, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
            
            p_v = cell_val.paragraphs[0]
            p_v.paragraph_format.line_spacing = 1.2
            if r_idx == 3: # Status row
                set_font_times(p_v.runs[0], size_pt=10.5, bold=True, color_rgb=RGBColor(0x00, 0x80, 0x00))
            else:
                set_font_times(p_v.runs[0], size_pt=10, color_rgb=RGBColor(0x33, 0x33, 0x33))
                
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ─── SECTION 4: HÌNH ẢNH MINH HỌA GIAO DIỆN KIỂM THỬ ───
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(4)
    h4.paragraph_format.keep_with_next = True
    r_h4 = h4.add_run("4. HÌNH ẢNH THỰC TẾ GIAO DIỆN KIỂM THỬ HỆ THỐNG")
    set_font_times(r_h4, size_pt=13.5, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    img_dir = r"C:\Users\ADMIN\Desktop\fire_detection_ai-main\docs\extracted_images"
    screenshots = [
        ("ui_overview.png", "Hình 4.1: Tổng quan giao diện giám sát FireEyes AI (11 thẻ số liệu CSDL và Camera AI)"),
        ("ui_camera_monitor.png", "Hình 4.2: Giao diện màn hình giám sát camera thời gian thực và thanh điều khiển"),
        ("ui_alarm_incident.png", "Hình 4.3: Trạng thái kích hoạt còi hú báo động đỏ khi phát hiện ngọn lửa bật lửa"),
        ("ui_alarm_detail.png", "Hình 4.4: Hộp thoại modal phóng to xem bằng chứng ảnh chụp snapshot thời điểm cháy"),
    ]
    
    for img_name, caption in screenshots:
        img_path = os.path.join(img_dir, img_name)
        if os.path.exists(img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(6)
            p_img.paragraph_format.space_after = Pt(2)
            doc.add_picture(img_path, width=Inches(5.8))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(8)
            r_cap = p_cap.add_run(caption)
            set_font_times(r_cap, size_pt=9.5, italic=True, color_rgb=RGBColor(0x55, 0x55, 0x55))
            
    # ─── SECTION 5: BẢNG TỔNG HỢP & HƯỚNG DẪN XỬ LÝ LỖI ───
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(4)
    h5.paragraph_format.keep_with_next = True
    r_h5 = h5.add_run("5. BẢNG TỔNG HỢP KẾT QUẢ VÀ HƯỚNG DẪN XỬ LÝ SỰ CỐ (TROUBLESHOOTING)")
    set_font_times(r_h5, size_pt=13.5, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    tbl_sum = doc.add_table(rows=11, cols=4)
    tbl_sum.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sum)
    
    sum_headers = ["Mã TC", "Tên ca kiểm thử", "Phân loại kiểm thử", "Kết quả thực tế"]
    for i, h in enumerate(sum_headers):
        cell = tbl_sum.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
        p = cell.paragraphs[0]
        set_font_times(p.runs[0], size_pt=10, bold=True, color_rgb=RGBColor(0xFF, 0xFF, 0xFF))
        
    sum_data = [
        ("TC-01", "Bật/Tắt Webcam & Phần cứng", "Giao diện & Thiết bị", "ĐẠT (PASS)"),
        ("TC-02", "Nhận diện ngọn lửa bật lửa thật", "AI & Real-time Vision", "ĐẠT (PASS)"),
        ("TC-03", "Bộ lọc chống báo động giả (da/tường)", "Thuật toán lọc vật lý", "ĐẠT (PASS)"),
        ("TC-04", "Tải ảnh phân tích hiện trường", "Xử lý ảnh tĩnh", "ĐẠT (PASS)"),
        ("TC-05", "Tạo dữ liệu giả lập MVC Seed Demo", "Toàn vẹn CSDL", "ĐẠT (PASS)"),
        ("TC-06", "Vòng đời sự cố (Acknowledge/Resolve)", "Quy trình nghiệp vụ", "ĐẠT (PASS)"),
        ("TC-07", "Xem ảnh snapshot bằng chứng & cache", "Hiển thị & Lưu trữ", "ĐẠT (PASS)"),
        ("TC-08", "Quản trị qua lệnh CLI Backend", "Dòng lệnh hệ thống", "ĐẠT (PASS)"),
        ("TC-09", "Bộ kiểm thử tự động 26 unit tests", "Kiểm thử hồi quy", "ĐẠT (26/26 PASS)"),
        ("TC-10", "Tài liệu hóa API Swagger UI", "Giao tiếp REST API", "ĐẠT (PASS)"),
    ]
    
    for row_idx, r_data in enumerate(sum_data, 1):
        bg = "F7F9FC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(r_data):
            cell = tbl_sum.cell(row_idx, col_idx)
            cell.text = text
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            if col_idx == 3:
                set_font_times(p.runs[0], size_pt=9.5, bold=True, color_rgb=RGBColor(0x00, 0x80, 0x00))
            else:
                set_font_times(p.runs[0], size_pt=9.5, color_rgb=RGBColor(0x33, 0x33, 0x33))
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    
    # Troubleshooting guide
    p_tb_title = doc.add_paragraph()
    p_tb_title.paragraph_format.space_before = Pt(6)
    p_tb_title.paragraph_format.space_after = Pt(4)
    r = p_tb_title.add_run("Hướng dẫn khắc phục sự cố thường gặp:")
    set_font_times(r, size_pt=11.5, bold=True, color_rgb=RGBColor(0x00, 0x33, 0x66))
    
    troubles = [
        ("Trình duyệt không mở được camera", "Nhấp vào biểu tượng ổ khóa/máy ảnh trên thanh địa chỉ URL của trình duyệt, chọn 'Cho phép' (Allow) quyền truy cập Camera, sau đó bấm F5 để tải lại trang."),
        ("Cổng 8000 bị xung đột hoặc đang bị chiếm", "Mở PowerShell chạy lệnh netstat -ano | findstr 8000 để tìm PID, sau đó chạy taskkill /F /PID <PID> để giải phóng cổng hoặc chỉ định cổng khác khi chạy uvicorn."),
        ("Bật lửa quá xa camera không báo cháy", "Đảm bảo đưa bật lửa cách camera từ 30cm - 60cm ở khu vực có đủ ánh sáng, giữ ngọn lửa cháy liên tục trên 1.5 giây để bộ đệm cửa sổ trượt 3/5 khung hình kịp tích lũy."),
        ("Không có tiếng còi hú khi báo cháy", "Trình duyệt hiện đại chặn âm thanh tự động phát (Autoplay Policy). Người dùng chỉ cần nhấp chuột bất kỳ vào trang web 1 lần để cấp quyền âm thanh cho Web Audio API."),
    ]
    
    for tb_h, tb_b in troubles:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(4)
        r_h = p.add_run(f"• {tb_h}: ")
        set_font_times(r_h, size_pt=10.5, bold=True, color_rgb=RGBColor(0xB2, 0x22, 0x22))
        
        r_b = p.add_run(tb_b)
        set_font_times(r_b, size_pt=10.5)
        
    doc.save(output_path)
    print(f"File {output_path} generated successfully with standardized Times New Roman font!")

if __name__ == "__main__":
    out_file = r"C:\Users\ADMIN\Desktop\huong_dan_test.docx"
    build_test_guide(out_file)
